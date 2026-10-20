"""Describe the shape of a report export without printing what is in it.

Run on each file the STAC report viewer exports (.xls, .doc, .pdf):

    python inspect_export.py path/to/export.xls

It answers the questions the reader has to be built around: what the file
really is underneath its extension, where the header row is, and whether the
hyperlinks to each case came through as usable URLs.

The export holds defendant names, dates of birth and case numbers, so by
default nothing from a data row is printed. The header row is printed because
it is column labels. A link is shown as its scheme, host and a masked path
shape, never the URL itself, since the path or query may carry a case id or a
token. The name and date of birth summary is counts only.

    python inspect_export.py --show-rows path/to/export.xls

prints every row with how the tool would read it, for Nick to review in his
own terminal. That output is criminal justice information: it is not for
pasting into a chat, a ticket or a commit.

Read-only. Writes nothing and logs nothing.
"""

import argparse
import logging
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import parse_qsl, urlsplit

from excel_handler import (
    COMMA_DELIMITER,
    MAX_SPLIT_PARAMETER,
    format_dob,
    parse_name,
    strip_suffix,
)

EXIT_OK = 0
EXIT_USAGE = 2
EXIT_CANNOT_READ = 1

# How many bytes to read to tell the formats apart.
SNIFF_BYTES = 512

# File signatures. The extension is not trusted: report servers commonly save
# an HTML table under an .xls name, and RTF under a .doc name.
SIG_ZIP = b"PK\x03\x04"
SIG_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
SIG_PDF = b"%PDF"
SIG_RTF = b"{\\rtf"

FORMAT_XLSX = "xlsx (zip)"
FORMAT_DOCX = "docx (zip)"
FORMAT_XLS = "binary xls (OLE, Excel 97-2003)"
FORMAT_DOC = "binary doc (OLE, Word 97-2003)"
FORMAT_OLE_UNKNOWN = "OLE compound file, not identified"
FORMAT_PDF = "pdf"
FORMAT_RTF = "rtf"
FORMAT_HTML = "html"
FORMAT_UNKNOWN = "unknown"

# The cell text that marks the header row. Rows above it are report titles
# and print dates; rows below it are defendants.
HEADER_MARKER = "UCN"

# How far down to look for the header row before giving up.
HEADER_SEARCH_ROWS = 50

# The shape of a UCN, e.g. 282026CF000000CFAXMX. Used only to answer yes or
# no to "does this link carry a case number", never to print one.
UCN_PATTERN = re.compile(r"\d{6}[A-Z]{2}\d{6}[A-Z]{2,8}", re.IGNORECASE)

# Path segments made only of letters are route names (Case, Details) and are
# safe to show. Anything else could be an id, so it is masked.
SAFE_SEGMENT = re.compile(r"^[A-Za-z_.-]+$")
MASKED = "<id>"
MASKED_VALUE = "<v>"

# Word and RTF write external links as a HYPERLINK field.
RTF_HYPERLINK = re.compile(rb'HYPERLINK\s+"([^"]+)"')

# Header labels for the columns the tool reads, matched case-insensitively.
NAME_HEADER = "NAME"
DOB_HEADER = "DOB"

# A DOB written as text in one of these is something the tool can convert.
# Anything else in a text cell would reach CCIS as typed.
DOB_TEXT_FORMATS = ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%m-%d-%Y")

SHOW_ROWS_BANNER = (
    "=" * 72
    + "\n  DEFENDANT DATA BELOW. For review in this terminal only.\n"
    "  Do not paste it into a chat, ticket, email or commit.\n"
    + "=" * 72
)
SHOW_ROWS_END = "=" * 72 + "\n  END OF DEFENDANT DATA\n" + "=" * 72


class CannotInspect(Exception):
    """The file is a format this script cannot look inside, or a library is missing."""


# --------------------------------------------------------------- helpers


def sniff(path: Path) -> str:
    """Work out the real format from the file's first bytes.

    Args:
        path: The export.

    Returns:
        One of the FORMAT_ constants.
    """
    head = path.read_bytes()[:SNIFF_BYTES]
    suffix = path.suffix.lower()

    if head.startswith(SIG_ZIP):
        try:
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
        except zipfile.BadZipFile:
            return FORMAT_UNKNOWN
        if any(name.startswith("xl/") for name in names):
            return FORMAT_XLSX
        if any(name.startswith("word/") for name in names):
            return FORMAT_DOCX
        return FORMAT_UNKNOWN

    if head.startswith(SIG_OLE):
        # Both old Excel and old Word are OLE containers. Telling them apart
        # properly means reading the directory, which the extension does well
        # enough for an inspection script.
        if suffix == ".xls":
            return FORMAT_XLS
        if suffix == ".doc":
            return FORMAT_DOC
        return FORMAT_OLE_UNKNOWN

    if head.startswith(SIG_PDF):
        return FORMAT_PDF

    text = head.lstrip(b"\xef\xbb\xbf").lstrip()
    if text.startswith(SIG_RTF):
        return FORMAT_RTF
    if text.startswith(b"<"):
        return FORMAT_HTML

    return FORMAT_UNKNOWN


def describe_link(url: str) -> str:
    """Reduce a URL to its scheme, host and a masked path shape.

    Args:
        url: A hyperlink target.

    Returns:
        e.g. 'https stac.example.com /Case/<id>?caseId=<v>', or the scheme
        alone for a javascript: or mailto: link.
    """
    parts = urlsplit(url.strip())
    if parts.scheme and not parts.netloc:
        # javascript:, mailto: and the like. Only the scheme says anything
        # useful, and the rest could be a call carrying a case id.
        return f"{parts.scheme}: (no host)"
    if not parts.netloc:
        return "relative link (no host)"

    segments = [
        segment if SAFE_SEGMENT.match(segment) else MASKED
        for segment in parts.path.split("/")
        if segment
    ]
    shape = "/" + "/".join(segments)
    params = [name for name, _ in parse_qsl(parts.query, keep_blank_values=True)]
    if params:
        shape += "?" + "&".join(f"{name}={MASKED_VALUE}" for name in params)
    return f"{parts.scheme} {parts.netloc} {shape}"


def report_links(links: list[str], ucns: list[str] | None = None) -> None:
    """Print a summary of a set of links, grouped by shape.

    Args:
        links: Hyperlink targets.
        ucns: The UCN from the same row as each link, when known. Used to
            count how many links carry their own row's case number.
    """
    print(f"  links found: {len(links)}")
    if not links:
        return

    shapes = Counter(describe_link(link) for link in links)
    for shape, count in shapes.most_common():
        print(f"    {count:>5}  {shape}")

    with_ucn_shape = sum(1 for link in links if UCN_PATTERN.search(link))
    print(f"  links containing something shaped like a UCN: {with_ucn_shape}")

    if ucns is not None:
        own = sum(
            1
            for link, ucn in zip(links, ucns)
            if ucn and ucn.upper() in link.upper()
        )
        print(f"  links containing their own row's UCN: {own}")


# ------------------------------------------------------------ table forms

# One cell: its value (str, or a datetime for a real date cell) and the link
# targets inside it.
Cell = tuple[object, list[str]]


def _text(value: object) -> str:
    """A cell value as text, for header matching and display."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def dob_for_ccis(value: object) -> tuple[str | None, str]:
    """What a DOB cell would become in the CCIS search, and what kind of cell it was.

    Args:
        value: The raw cell value.

    Returns:
        (mm/dd/yyyy or None, a short label for the cell kind). A text cell in
        an unrecognised layout comes back as typed, labelled so, because that
        is what would reach CCIS.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "blank"
    if isinstance(value, (datetime, date)):
        return format_dob(value), "date cell"
    if isinstance(value, str):
        for layout in DOB_TEXT_FORMATS:
            try:
                return datetime.strptime(value.strip(), layout).strftime("%m/%d/%Y"), "text"
            except ValueError:
                continue
        return value.strip(), "text, unrecognised layout"
    return None, f"{type(value).__name__}, not a date"


def name_notes(name: str) -> list[str]:
    """Labels for anything about a name that changes how it is searched.

    Args:
        name: The name cell, 'Last, First Middle'.

    Returns:
        Zero or more short labels. Empty means a plain 'Last, First'.
    """
    if not name.strip():
        return ["blank"]
    if name.count(COMMA_DELIMITER) > MAX_SPLIT_PARAMETER:
        return ["more than one comma (row would be refused)"]
    first, last = parse_name(name)
    if not first or not last:
        return ["cannot split into Last, First (row would be refused)"]

    notes = []
    search_last = strip_suffix(last)
    if search_last != last:
        notes.append("suffix dropped for CCIS")
    if len(search_last.split()) > 1:
        notes.append("compound surname")
    if "-" in search_last:
        notes.append("hyphenated surname")
    if len(name.split(COMMA_DELIMITER, 1)[1].split()) > 1:
        notes.append("middle name or initial (not searched)")
    return notes


def _column(header: list[str], label: str) -> int | None:
    """Index of the header cell reading `label`, ignoring case. None if absent."""
    for index, text in enumerate(header):
        if text.upper() == label:
            return index
    return None


def _cell(row: list[Cell], index: int | None) -> Cell:
    """The cell at `index`, or an empty one when the row is short or the column absent."""
    if index is None or index >= len(row):
        return None, []
    return row[index]


def report_names_and_dobs(header: list[str], data: list[list[Cell]], show_rows: bool) -> None:
    """Summarise how the names and DOBs would be read, and optionally list every row.

    Args:
        header: The header labels.
        data: The rows below the header.
        show_rows: Print each row's values. Defendant data, for Nick's
            terminal only.
    """
    ucn_at = _column(header, HEADER_MARKER)
    name_at = _column(header, NAME_HEADER)
    dob_at = _column(header, DOB_HEADER)

    if name_at is None or dob_at is None:
        print(f"  no {NAME_HEADER!r} and/or {DOB_HEADER!r} column, skipping the name check")
        return

    name_counts: Counter = Counter()
    dob_counts: Counter = Counter()
    listing = []
    for number, row in enumerate(data, start=1):
        name = _text(_cell(row, name_at)[0]).strip()
        dob_sent, dob_kind = dob_for_ccis(_cell(row, dob_at)[0])
        notes = name_notes(name)
        name_counts.update(notes or ["plain Last, First"])
        dob_counts[dob_kind] += 1

        if show_rows:
            first, last = parse_name(name)
            ucn_value, ucn_links = _cell(row, ucn_at)
            listing.append(
                f"  {number:>3}  {_text(ucn_value):<22} {name}\n"
                f"       CCIS search: first {first!r}, last "
                f"{strip_suffix(last) if last else None!r}, dob {dob_sent!r} ({dob_kind})\n"
                f"       link: {'yes' if ucn_links else 'NO'}"
                + (f"   notes: {', '.join(notes)}" if notes else "")
            )

    print("  names:")
    for label, count in name_counts.most_common():
        print(f"    {count:>5}  {label}")
    print("  dates of birth:")
    for label, count in dob_counts.most_common():
        print(f"    {count:>5}  {label}")

    if show_rows:
        print(SHOW_ROWS_BANNER)
        for entry in listing:
            print(entry)
        print(SHOW_ROWS_END)


def report_table(rows: list[list[Cell]], show_rows: bool = False) -> None:
    """Describe a table: header row, sizes, the links in each column, names and DOBs.

    Shared by xls, xlsx and html, which all come down to rows of cells.

    Args:
        rows: Each row a list of (cell value, [link targets in that cell]).
        show_rows: Also print every data row. See report_names_and_dobs.
    """
    header_at = None
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        if any(_text(value).strip().upper() == HEADER_MARKER for value, _ in row):
            header_at = index
            break

    print(f"  rows in file: {len(rows)}")
    if header_at is None:
        print(
            f"  no header row found (no cell reading {HEADER_MARKER!r} in the first "
            f"{HEADER_SEARCH_ROWS} rows). Not printing any row, since one may be data."
        )
        return

    header = [_text(value).strip() for value, _ in rows[header_at]]
    ucn_column = _column(header, HEADER_MARKER)
    data = [
        row
        for row in rows[header_at + 1:]
        if any(_text(value).strip() for value, _ in row)
    ]

    print(f"  header is row {header_at + 1} (rows above it: {header_at})")
    print(f"  columns: {len(header)}")
    for number, label in enumerate(header, start=1):
        print(f"    {number:>2}  {label or '(blank)'}")
    print(f"  non-empty rows below the header: {len(data)}")

    links_by_column: dict[int, list[tuple[str, str]]] = defaultdict(list)
    for row in data:
        ucn = _text(_cell(row, ucn_column)[0]).strip()
        for column, (_, targets) in enumerate(row):
            for target in targets:
                links_by_column[column].append((target, ucn))

    if not links_by_column:
        print("  no hyperlinks in any data row")
    for column in sorted(links_by_column):
        label = header[column] if column < len(header) else "(beyond header)"
        pairs = links_by_column[column]
        print(f"  column {column + 1} ({label}):")
        report_links([link for link, _ in pairs], [ucn for _, ucn in pairs])

    report_names_and_dobs(header, data, show_rows)


def read_xls(path: Path) -> list[list[Cell]]:
    """Read a binary .xls into rows of (text, links).

    Raises:
        CannotInspect: xlrd is not installed.
    """
    try:
        import xlrd
    except ImportError as error:
        raise CannotInspect(
            "This is a binary .xls and reading it needs xlrd. Run: pip install xlrd"
        ) from error

    book = xlrd.open_workbook(str(path), on_demand=True)
    print(f"  sheets: {book.sheet_names()}")
    sheet = book.sheet_by_index(0)
    links = sheet.hyperlink_map

    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            # xlrd hands a date cell back as a day count. Converting it here
            # is what lets the DOB summary tell a real date from typed text.
            if cell.ctype == xlrd.XL_CELL_DATE:
                value: object = xlrd.xldate_as_datetime(cell.value, book.datemode)
            else:
                value = cell.value
            link = links.get((r, c))
            row.append((value, [link.url_or_path] if link and link.url_or_path else []))
        rows.append(row)
    book.release_resources()
    return rows


def read_xlsx(path: Path) -> list[list[Cell]]:
    """Read an .xlsx into rows of (text, links), first sheet only."""
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=False)
    try:
        print(f"  sheets: {book.sheetnames}")
        sheet = book.worksheets[0]
        rows = []
        for cells in sheet.iter_rows():
            rows.append([
                (
                    cell.value,
                    [cell.hyperlink.target] if cell.hyperlink and cell.hyperlink.target else [],
                )
                for cell in cells
            ])
        return rows
    finally:
        book.close()


class _TableParser(HTMLParser):
    """Collects <tr> rows of <td>/<th> cells, with the hrefs inside each cell."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables = 0
        self.rows: list[list[tuple[str, list[str]]]] = []
        self.loose_links: list[str] = []
        self._row: list[tuple[str, list[str]]] | None = None
        self._text: list[str] | None = None
        self._links: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        """Open rows and cells, and record links."""
        if tag == "table":
            self.tables += 1
        elif tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._text, self._links = [], []
        elif tag == "a":
            href = dict(attrs).get("href")
            if href:
                if self._links is not None:
                    self._links.append(href)
                else:
                    self.loose_links.append(href)

    def handle_endtag(self, tag: str) -> None:
        """Close cells and rows."""
        if tag in ("td", "th") and self._row is not None and self._text is not None:
            self._row.append((" ".join("".join(self._text).split()), self._links or []))
            self._text, self._links = None, None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        """Gather cell text."""
        if self._text is not None:
            self._text.append(data)


def read_html(path: Path) -> list[list[Cell]]:
    """Read an HTML table (often saved as .xls or .doc) into rows of (text, links)."""
    parser = _TableParser()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    print(f"  <table> elements: {parser.tables}")
    if parser.loose_links:
        print("  links outside any table cell:")
        report_links(parser.loose_links)
    return parser.rows


# ------------------------------------------------------------- link only


def inspect_pdf(path: Path) -> None:
    """Report the link annotations in a PDF.

    Raises:
        CannotInspect: pypdf is not installed.
    """
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise CannotInspect(
            "Reading PDF links needs pypdf. Run: pip install pypdf"
        ) from error

    reader = PdfReader(str(path))
    print(f"  pages: {len(reader.pages)}")
    links = []
    for page in reader.pages:
        for annotation in page.get("/Annots") or []:
            action = annotation.get_object().get("/A") or {}
            uri = action.get("/URI")
            if uri:
                links.append(str(uri))
    report_links(links)


def inspect_rtf(path: Path) -> None:
    """Report the HYPERLINK fields in an RTF file."""
    links = [match.decode("latin-1") for match in RTF_HYPERLINK.findall(path.read_bytes())]
    report_links(links)


def inspect_docx(path: Path) -> None:
    """Report the external links in a .docx, from its relationships file."""
    with zipfile.ZipFile(path) as archive:
        try:
            rels = archive.read("word/_rels/document.xml.rels").decode("utf-8", "replace")
        except KeyError:
            rels = ""
    links = re.findall(r'Target="([^"]+)"[^>]*TargetMode="External"', rels)
    report_links(links)


# ------------------------------------------------------------------ main


def inspect(path: Path, show_rows: bool = False) -> None:
    """Print the shape of one export.

    Args:
        path: The export file.
        show_rows: Also print every data row of a table format.

    Raises:
        CannotInspect: The format cannot be looked inside here.
    """
    kind = sniff(path)
    print(f"{path.name}")
    print(f"  extension: {path.suffix or '(none)'}   real format: {kind}")
    print(f"  size: {path.stat().st_size // 1024} KB")

    if kind == FORMAT_XLS:
        report_table(read_xls(path), show_rows)
    elif kind == FORMAT_XLSX:
        report_table(read_xlsx(path), show_rows)
    elif kind == FORMAT_HTML:
        report_table(read_html(path), show_rows)
    elif kind == FORMAT_PDF:
        inspect_pdf(path)
    elif kind == FORMAT_RTF:
        inspect_rtf(path)
    elif kind == FORMAT_DOCX:
        inspect_docx(path)
    else:
        raise CannotInspect(
            f"Cannot look inside a {kind} file here. The format above is the answer "
            "for this one."
        )


def main(argv: list[str]) -> int:
    """Inspect each file named on the command line.

    Args:
        argv: File paths, optionally with --show-rows.

    Returns:
        EXIT_OK if every file was inspected, EXIT_CANNOT_READ if any was not,
        EXIT_USAGE if none were given.
    """
    parser = argparse.ArgumentParser(
        description="Describe a report export's shape without printing its contents."
    )
    parser.add_argument("files", nargs="*", help="Export files to inspect.")
    parser.add_argument(
        "--show-rows",
        action="store_true",
        help="Also print every row and how it would be searched. Defendant data: "
             "for review in this terminal only.",
    )
    args = parser.parse_args(argv)

    if not args.files:
        parser.print_usage(sys.stderr)
        return EXIT_USAGE

    # Crystal writes its PDF link entries with a duplicated key, and pypdf
    # warns once per link. Harmless, and it buries the output.
    logging.getLogger("pypdf").setLevel(logging.ERROR)

    status = EXIT_OK
    for name in args.files:
        path = Path(name)
        if not path.is_file():
            print(f"{name}: no such file", file=sys.stderr)
            status = EXIT_CANNOT_READ
            continue
        try:
            inspect(path, args.show_rows)
        except CannotInspect as error:
            print(f"  stopped: {error}")
            status = EXIT_CANNOT_READ
        except Exception as error:  # noqa: BLE001 - report the type, not the contents
            # Only the error type, since some library messages quote cell text.
            print(f"  stopped: could not read the file ({type(error).__name__})")
            status = EXIT_CANNOT_READ
        print()
    return status


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
