"""Unit tests for report_reader, using made-up rows only."""

import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import report_reader as rr  # noqa: E402
from excel_handler import Outcome, Spreadsheet  # noqa: E402
from exceptions import ConfigProblem  # noqa: E402

LINK = "https://stac-test.example.com/incidents/00000"
UCN = "282026CF000000CFAXMX"
OTHER_UCN = "282026CF000001CFAXMX"


def export(*data_rows: list) -> list[list]:
    """A made-up export in the real layout: title row, header, then data.

    Each data row is (ucn, link, name, dob); spacer columns are added here.
    """
    rows = [
        [("Made-up report title", ""), ("", ""), ("", ""), ("", ""), ("", ""), ("", "")],
        [("", ""), ("UCN", ""), ("OPEN Date", ""), ("Name", ""), ("", ""), ("DOB", "")],
    ]
    for ucn, link, name, dob in data_rows:
        rows.append([
            ("", ""), (ucn, link), ("01/01/2026", ""), (name, ""), ("", ""), (dob, ""),
        ])
    return rows


def build(rows: list[list]) -> list[dict]:
    header_at, columns = rr.find_header(rows)
    return rr.build_rows(rows, header_at, columns)


class HeaderTests(unittest.TestCase):
    def test_finds_header_below_title_and_skips_spacers(self):
        header_at, columns = rr.find_header(export())
        self.assertEqual(header_at, 1)
        self.assertEqual(columns["UCN"], 1)
        self.assertEqual(columns["NAME"], 3)
        self.assertEqual(columns["DOB"], 5)
        self.assertEqual(columns["OPEN DATE"], 2)

    def test_header_matching_ignores_case_and_spacing(self):
        rows = [[("ucn", ""), ("  name ", ""), ("Dob", "")]]
        _, columns = rr.find_header(rows)
        self.assertEqual(columns, {"UCN": 0, "NAME": 1, "DOB": 2})

    def test_missing_column_names_what_is_there(self):
        rows = [[("UCN", ""), ("Defendant", ""), ("DOB", "")]]
        with self.assertRaises(ConfigProblem) as caught:
            rr.find_header(rows)
        self.assertIn("NAME", str(caught.exception))
        self.assertIn("DEFENDANT", str(caught.exception))

    def test_no_header_at_all(self):
        with self.assertRaises(ConfigProblem):
            rr.find_header([[("nothing", "")]])


class DobTests(unittest.TestCase):
    def test_text_in_the_export_layout(self):
        self.assertEqual(rr.normalise_dob("1980-01-02"), "01/02/1980")
        self.assertEqual(rr.normalise_dob(" 1980-12-31 "), "12/31/1980")

    def test_month_first_with_dashes(self):
        self.assertEqual(rr.normalise_dob("01-01-1980"), "01/01/1980")
        self.assertEqual(rr.normalise_dob(" 1-2-1980 "), "01/02/1980")

    def test_slashes_are_accepted_too(self):
        self.assertEqual(rr.normalise_dob("01/01/1980"), "01/01/1980")
        self.assertEqual(rr.normalise_dob(" 1/2/1980 "), "01/02/1980")

    def test_real_date_cell(self):
        self.assertEqual(rr.normalise_dob(datetime(1980, 1, 1)), "01/01/1980")

    def test_other_layouts_are_refused(self):
        for value in (
            "1/1/80", "1-1-80", "13/01/1980", "13-01-1980", "02-30-1980",
            "1980-13-01", "1980/01/01", "sometime", "", None, 29221.0,
        ):
            self.assertIsNone(rr.normalise_dob(value), value)


class BuildRowsTests(unittest.TestCase):
    def test_plain_row(self):
        (row,) = build(export((UCN, LINK, "DOE, John", "01/01/1980")))
        self.assertEqual(row["row"], 2)
        self.assertEqual(row["export_row"], 3)
        self.assertEqual(row["first_name"], "John")
        self.assertEqual(row["last_name"], "DOE")
        self.assertEqual(row["search_last"], "DOE")
        self.assertEqual(row["dob"], "01/01/1980")
        self.assertEqual(row["ucn"], UCN)
        self.assertEqual(row["case_url"], LINK)
        self.assertEqual(row["case_number"], "")
        self.assertEqual(row["problem"], "")

    def test_suffix_is_dropped_for_the_search_only(self):
        (row,) = build(export((UCN, LINK, "DOE JR, John Q", "01/01/1980")))
        self.assertEqual(row["last_name"], "DOE JR")
        self.assertEqual(row["search_last"], "DOE")
        self.assertEqual(row["first_name"], "John")

    def test_missing_link_is_not_a_problem(self):
        (row,) = build(export((UCN, "", "DOE, John", "01/01/1980")))
        self.assertEqual(row["case_url"], "")
        self.assertEqual(row["problem"], "")

    def test_problem_rows_are_kept_and_flagged(self):
        rows = build(export(
            (UCN, LINK, "", "01/01/1980"),
            (UCN, LINK, "DOE, JR, John", "01/01/1980"),
            (UCN, LINK, "John Doe", "01/01/1980"),
            (UCN, LINK, "DOE, John", "1980-13-01"),
            ("", LINK, "DOE, John", "01/01/1980"),
        ))
        self.assertEqual(len(rows), 5)
        problems = [row["problem"] for row in rows]
        self.assertIn("no name", problems[0])
        self.assertIn("more than one comma", problems[1])
        self.assertIn("could not split", problems[2])
        self.assertIn("date of birth", problems[3])
        self.assertIn("no UCN", problems[4])

    def test_problems_never_quote_the_name_or_date(self):
        # Problems are logged, and the log never carries names or dates of
        # birth. The unreadable date matters most: the redactor cannot catch
        # it, because it never became a date of birth to register.
        rows = build(export(
            (UCN, LINK, "DOE, JR, John", "01/01/1980"),
            (UCN, LINK, "John Doe", "01/01/1980"),
            (UCN, LINK, "DOE, John", "1980-13-01"),
        ))
        for row in rows:
            text = row["problem"].upper()
            self.assertTrue(text)
            for value in ("DOE", "JOHN", "1980"):
                self.assertNotIn(value, text)

    def test_blank_rows_are_skipped_and_numbering_stays_dense(self):
        rows = export((UCN, LINK, "DOE, John", "01/01/1980"))
        rows.append([("", "")] * 6)
        rows += export((OTHER_UCN, LINK, "DOE, Jane", "02/02/1981"))[2:]
        built = build(rows)
        self.assertEqual([row["row"] for row in built], [2, 3])
        self.assertEqual([row["export_row"] for row in built], [3, 5])


class WorkingCopyTests(unittest.TestCase):
    def setUp(self):
        self._folder = tempfile.TemporaryDirectory()
        self.folder = Path(self._folder.name)
        self.path = self.folder / "AutoCCIS.xlsx"

    def tearDown(self):
        self._folder.cleanup()

    def test_written_with_link_and_readable_by_spreadsheet(self):
        rows = build(export(
            (UCN, LINK, "DOE, John", "01/01/1980"),
            (OTHER_UCN, "", "DOE, Jane", "02/02/1981"),
        ))
        rr.write_working_copy(self.path, rows)

        book = openpyxl.load_workbook(self.path)
        sheet = book[rr.WORKING_SHEET]
        self.assertEqual([c.value for c in sheet[1]], list(rr.WORKING_HEADERS))
        self.assertEqual(sheet.cell(row=2, column=1).hyperlink.target, LINK)
        self.assertIsNone(sheet.cell(row=3, column=1).hyperlink)
        self.assertEqual(sheet.cell(row=2, column=4).value, "01/01/1980")
        book.close()

        with Spreadsheet(rr.working_copy_config(self.path)) as sheet:
            loaded = sheet.load_rows()
            self.assertEqual([row["ucn"] for row in loaded], [UCN, OTHER_UCN])
            sheet.write_outcome(2, Outcome.FILED)

        self.assertEqual(rr.read_previous_outcomes(self.path), {UCN: Outcome.FILED})

    def test_outcomes_carry_over_by_ucn_on_a_restart(self):
        first = build(export(
            (UCN, LINK, "DOE, John", "01/01/1980"),
            (OTHER_UCN, LINK, "DOE, Jane", "02/02/1981"),
        ))
        first[0]["outcome"] = Outcome.CHECK_BY_HAND
        rr.write_working_copy(self.path, first)

        previous = rr.read_previous_outcomes(self.path)
        self.assertEqual(previous, {UCN: Outcome.CHECK_BY_HAND})

    def test_no_working_copy_means_no_outcomes(self):
        self.assertEqual(rr.read_previous_outcomes(self.path), {})


class RetargetLinkTests(unittest.TestCase):
    PROD = "https://stac-prod.example.org"
    TEST = "https://stac-test.example.com"

    def test_prod_link_is_pointed_at_test_with_path_kept(self):
        moved = rr.retarget_link(
            "https://stac-prod.example.org/incidents/00000?tab=3", self.PROD, self.TEST
        )
        self.assertEqual(moved, "https://stac-test.example.com/incidents/00000?tab=3")

    def test_host_match_ignores_case(self):
        moved = rr.retarget_link(
            "https://STAC-PROD.example.org/incidents/00000", self.PROD, self.TEST
        )
        self.assertEqual(moved, "https://stac-test.example.com/incidents/00000")

    def test_other_hosts_and_blanks_are_left_alone(self):
        for url in ("", "https://elsewhere.example.net/incidents/1", "javascript:go(1)"):
            self.assertEqual(rr.retarget_link(url, self.PROD, self.TEST), url)

    def test_same_instance_changes_nothing(self):
        url = "https://stac-test.example.com/incidents/00000"
        self.assertEqual(rr.retarget_link(url, self.TEST, self.TEST), url)


class ReadXlsTests(unittest.TestCase):
    def test_refuses_a_file_that_is_not_binary_xls(self):
        with tempfile.TemporaryDirectory() as folder:
            fake = Path(folder) / "AutoCCIS.xls"
            fake.write_text("<html><table></table></html>")
            with self.assertRaises(ConfigProblem) as caught:
                rr.read_xls(fake)
            self.assertIn("inspect_export.py", str(caught.exception))

    def test_missing_file(self):
        with self.assertRaises(ConfigProblem):
            rr.read_xls(Path("no-such-export.xls"))


if __name__ == "__main__":
    unittest.main()
